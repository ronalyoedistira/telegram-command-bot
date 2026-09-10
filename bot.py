import logging
import os
import tempfile
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from telegram import InlineKeyboardButton, InlineKeyboardMarkup, Update
from telegram.ext import (
    ApplicationBuilder,
    CallbackQueryHandler,
    CommandHandler,
    ContextTypes,
)

import database

TOKEN = os.getenv("BOT_TOKEN")


def format_rupiah(angka: int) -> str:
    """Ubah angka jadi format 'Rp 50.000'."""
    return f"Rp {angka:,.0f}".replace(",", ".")


def buat_file_export(user_id: int, transaksi: list) -> str:
    """Buat file .xlsx berisi seluruh transaksi milik user_id. Mengembalikan path file sementara."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Transaksi"

    header = ["ID", "Tanggal", "Jenis", "Kategori", "Jumlah (Rp)", "Keterangan"]
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="305496")

    total_masuk = 0
    total_keluar = 0
    for t in transaksi:
        tanggal = t["waktu"][:16].replace("T", " ")
        jenis_label = "Masuk" if t["jenis"] == "masuk" else "Keluar"
        ws.append([t["id"], tanggal, jenis_label, t["kategori"], t["jumlah"], t["keterangan"]])
        if t["jenis"] == "masuk":
            total_masuk += t["jumlah"]
        else:
            total_keluar += t["jumlah"]

    ws.append([])
    ws.append(["", "", "", "", "TOTAL MASUK", total_masuk])
    ws.append(["", "", "", "", "TOTAL KELUAR", total_keluar])
    ws.append(["", "", "", "", "SALDO", total_masuk - total_keluar])
    for row in ws.iter_rows(min_row=ws.max_row - 2, max_row=ws.max_row):
        for cell in row:
            cell.font = Font(bold=True)

    for col, width in zip("ABCDEF", [6, 18, 10, 16, 14, 30]):
        ws.column_dimensions[col].width = width

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    path_file = os.path.join(tempfile.gettempdir(), f"export_transaksi_{user_id}_{timestamp}.xlsx")
    wb.save(path_file)
    return path_file


logging.basicConfig(
    format="%(asctime)s - %(name)s - %(levelname)s - %(message)s",
    level=logging.INFO,
)

# ============================================================
# TEKS & KONTEN
# ============================================================

MENU_TEXT = (
    "🤖 *MENU UTAMA*\n\n"
    "Silakan pilih fitur yang tersedia di bawah ini.\n\n"
    "Bot ini dirancang untuk merespons command dan tombol menu tertentu."
)

INFO_MENU_TEXT = (
    "ℹ️ *INFO & REFERENSI*\n\n"
    "Pilih topik di bawah untuk melihat penjelasan lengkapnya.\n"
    "Menu ini berisi rangkuman istilah, FAQ, dan tips yang sering dicari pengguna bot."
)

# Konten referensi asli (bukan placeholder) — mudah ditambah/diedit di satu tempat ini.
INFO_CONTENT = {
    "info_istilah": {
        "title": "📖 ISTILAH TEKNIS",
        "text": (
            "📖 *ISTILAH TEKNIS*\n\n"
            "*Command*\nPerintah yang diawali garis miring, contoh `/start`. Dikenali otomatis oleh Telegram sebagai aksi khusus, bukan pesan teks biasa.\n\n"
            "*Inline Keyboard*\nTombol yang muncul menempel di bawah sebuah pesan (bukan di keyboard HP). Setiap tombol punya `callback_data` yang dikirim balik ke bot saat ditekan.\n\n"
            "*Callback Query*\nData yang dikirim ke bot saat pengguna menekan tombol inline. Bot wajib membalas dengan `answer()` agar tombol tidak terlihat 'loading' terus.\n\n"
            "*Polling*\nCara bot mengambil pesan dengan terus-menerus bertanya ke server Telegram \"ada pesan baru?\". Tidak butuh URL publik/ngrok, tapi proses botnya harus terus menyala.\n\n"
            "*Webhook*\nKebalikan dari polling — Telegram yang mengirim pesan ke URL bot secara otomatis. Butuh server dengan URL publik (misalnya lewat ngrok saat development)."
        ),
    },
    "info_faq": {
        "title": "❓ FAQ",
        "text": (
            "❓ *FAQ (PERTANYAAN UMUM)*\n\n"
            "*Q: Kenapa bot tidak merespons sama sekali?*\n"
            "A: Cek apakah proses `python bot.py` masih berjalan di terminal, dan pastikan `BOT_TOKEN` sudah benar diisi.\n\n"
            "*Q: Bot merespons command tapi tombol tidak bereaksi?*\n"
            "A: Pastikan `CallbackQueryHandler` sudah didaftarkan di fungsi `main()`.\n\n"
            "*Q: Apakah bot ini butuh ngrok?*\n"
            "A: Tidak. Bot ini pakai *polling*, jadi bisa langsung jalan dari laptop tanpa URL publik.\n\n"
            "*Q: Bagaimana kalau saya ingin ganti/reset token?*\n"
            "A: Buka chat dengan @BotFather, kirim `/revoke`, pilih bot-nya, lalu salin token baru yang diberikan.\n\n"
            "*Q: Bisakah bot ini dijalankan 24 jam tanpa laptop menyala?*\n"
            "A: Bisa, dengan deploy ke layanan hosting seperti Railway (lihat README.md di project ini)."
        ),
    },
    "info_keamanan": {
        "title": "🔒 TIPS KEAMANAN",
        "text": (
            "🔒 *TIPS KEAMANAN*\n\n"
            "1️⃣ Jangan pernah menulis token langsung di dalam kode `bot.py` — gunakan environment variable (`BOT_TOKEN`).\n\n"
            "2️⃣ Jangan commit file `.env` ke GitHub. Pastikan `.gitignore` sudah mengecualikannya.\n\n"
            "3️⃣ Kalau token pernah tidak sengaja terunggah ke tempat publik, langsung `/revoke` lewat @BotFather untuk membatalkannya.\n\n"
            "4️⃣ Batasi siapa yang bisa memakai command sensitif (misalnya command admin) dengan mengecek `update.effective_user.id`.\n\n"
            "5️⃣ Untuk bot produksi, simpan token di fitur *Variables/Secrets* platform hosting (Railway, dst), bukan di file biasa."
        ),
    },
    "info_tentang": {
        "title": "🚀 TENTANG BOT INI",
        "text": (
            "🚀 *TENTANG BOT INI*\n\n"
            "Bot ini dibangun dengan library `python-telegram-bot`, memakai arsitektur *polling* (bot yang aktif bertanya ke Telegram, bukan menunggu dikirimi).\n\n"
            "*Command yang tersedia:*\n"
            "`/start` — pesan sambutan + menu utama\n"
            "`/menu` — buka menu utama kapan saja\n"
            "`/help` — bantuan & daftar command\n"
            "`/info` — menu info & referensi (yang sedang kamu buka ini)\n\n"
            "Struktur project ini didesain agar mudah dikembangkan lebih lanjut — misalnya ditambah database, integrasi payment gateway, atau webhook eksternal."
        ),
    },
}


# ============================================================
# KEYBOARD (MENU TOMBOL)
# ============================================================

def main_menu_keyboard() -> InlineKeyboardMarkup:
    keyboard = [
        [
            InlineKeyboardButton("📊 Laporan", callback_data="menu_laporan"),
            InlineKeyboardButton("💰 Saldo", callback_data="menu_saldo"),
        ],
        [
            InlineKeyboardButton("🧾 Transaksi", callback_data="menu_transaksi"),
            InlineKeyboardButton("📁 Data", callback_data="menu_data"),
        ],
        [
            InlineKeyboardButton("ℹ️ Info & Referensi", callback_data="menu_info"),
            InlineKeyboardButton("❓ Bantuan", callback_data="menu_help"),
        ],
    ]
    return InlineKeyboardMarkup(keyboard)


def help_keyboard() -> InlineKeyboardMarkup:
    keyboard = [
        [InlineKeyboardButton("📖 Cara Menggunakan", callback_data="help_cara")],
        [InlineKeyboardButton("📋 Daftar Command", callback_data="help_command")],
        [InlineKeyboardButton("ℹ️ Info & Referensi", callback_data="menu_info")],
        [InlineKeyboardButton("🏠 Menu Utama", callback_data="back_menu")],
    ]
    return InlineKeyboardMarkup(keyboard)


def info_menu_keyboard() -> InlineKeyboardMarkup:
    keyboard = [
        [InlineKeyboardButton("📖 Istilah Teknis", callback_data="info_istilah")],
        [InlineKeyboardButton("❓ FAQ", callback_data="info_faq")],
        [InlineKeyboardButton("🔒 Tips Keamanan", callback_data="info_keamanan")],
        [InlineKeyboardButton("🚀 Tentang Bot Ini", callback_data="info_tentang")],
        [InlineKeyboardButton("🏠 Menu Utama", callback_data="back_menu")],
    ]
    return InlineKeyboardMarkup(keyboard)


def info_detail_keyboard() -> InlineKeyboardMarkup:
    keyboard = [
        [InlineKeyboardButton("⬅️ Kembali ke Info", callback_data="menu_info")],
        [InlineKeyboardButton("🏠 Menu Utama", callback_data="back_menu")],
    ]
    return InlineKeyboardMarkup(keyboard)


def back_menu_keyboard() -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup(
        [[InlineKeyboardButton("🏠 Menu Utama", callback_data="back_menu")]]
    )


# ============================================================
# HANDLER COMMAND (/start, /help, /menu, /info)
# ============================================================

async def start(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    text = (
        "👋 *Selamat datang!*\n\n"
        "Saya adalah Telegram Bot Anda.\n\n"
        "Saya dapat membantu Anda melalui command dan menu interaktif — "
        "termasuk menu *Info & Referensi* yang berisi penjelasan istilah, FAQ, dan tips.\n\n"
        "Gunakan /menu untuk melihat fitur, /info untuk referensi, atau /help untuk bantuan."
    )
    await update.message.reply_text(text, parse_mode="Markdown", reply_markup=main_menu_keyboard())


async def help_command(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    text = (
        "❓ *BANTUAN*\n\n"
        "Command yang tersedia:\n\n"
        "/start — Memulai bot\n"
        "/menu — Membuka menu utama\n"
        "/info — Membuka menu info & referensi\n"
        "/help — Menampilkan bantuan\n\n"
        "Anda juga dapat menggunakan tombol pada menu untuk berpindah fitur tanpa mengetik command."
    )
    await update.message.reply_text(text, parse_mode="Markdown", reply_markup=help_keyboard())


async def menu(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    await update.message.reply_text(
        MENU_TEXT,
        parse_mode="Markdown",
        reply_markup=main_menu_keyboard(),
    )


async def info_command(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    """Command /info — membuka menu referensi langsung dari chat (tanpa lewat /menu dulu)."""
    await update.message.reply_text(
        INFO_MENU_TEXT,
        parse_mode="Markdown",
        reply_markup=info_menu_keyboard(),
    )


# ============================================================
# HANDLER COMMAND TRANSAKSI (/masuk, /keluar, /reset_data)
# ============================================================

async def catat_masuk(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    """Format: /masuk 50000 Gaji bulanan #Gaji  (bagian #Kategori opsional, default 'Umum')"""
    await _catat_transaksi(update, context, jenis="masuk")


async def catat_keluar(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    """Format: /keluar 20000 Makan siang #Makanan  (bagian #Kategori opsional, default 'Umum')"""
    await _catat_transaksi(update, context, jenis="keluar")


def _ekstrak_kategori(args: list) -> tuple:
    """Cari token yang diawali '#' di antara args (contoh '#Makanan'), pisahkan dari keterangan.
    Mengembalikan (kategori, args_tanpa_tag_kategori). Kalau tidak ada tag, kategori = None."""
    kategori = None
    sisa_args = []
    for a in args:
        if a.startswith("#") and len(a) > 1 and kategori is None:
            kategori = a[1:]
        else:
            sisa_args.append(a)
    return kategori, sisa_args


async def _catat_transaksi(update: Update, context: ContextTypes.DEFAULT_TYPE, jenis: str) -> None:
    contoh = (
        "/masuk 50000 Gaji bulanan #Gaji" if jenis == "masuk"
        else "/keluar 20000 Makan siang #Makanan"
    )

    if not context.args:
        await update.message.reply_text(
            f"⚠️ Format salah.\n\nContoh penggunaan:\n`{contoh}`\n\n"
            "Bagian `#Kategori` di akhir bersifat opsional — kalau tidak diisi, otomatis masuk kategori 'Umum'.",
            parse_mode="Markdown",
        )
        return

    jumlah_str = context.args[0]

    if not jumlah_str.isdigit():
        await update.message.reply_text(
            f"⚠️ Jumlah harus berupa angka bulat (tanpa titik/koma).\n\nContoh:\n`{contoh}`",
            parse_mode="Markdown",
        )
        return

    kategori, sisa_args = _ekstrak_kategori(context.args[1:])
    keterangan = " ".join(sisa_args) if sisa_args else "-"
    kategori = kategori or database.KATEGORI_DEFAULT

    jumlah = int(jumlah_str)
    user_id = update.effective_user.id
    database.tambah_transaksi(user_id, jenis, jumlah, keterangan, kategori=kategori)
    saldo_baru = database.get_saldo(user_id)

    ikon = "🟢" if jenis == "masuk" else "🔴"
    label = "Pemasukan" if jenis == "masuk" else "Pengeluaran"

    await update.message.reply_text(
        f"{ikon} *{label} tercatat*\n\n"
        f"Jumlah: {format_rupiah(jumlah)}\n"
        f"Kategori: {kategori}\n"
        f"Keterangan: {keterangan}\n\n"
        f"💰 Saldo sekarang: {format_rupiah(saldo_baru)}",
        parse_mode="Markdown",
    )


async def reset_data(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    user_id = update.effective_user.id
    jumlah_dihapus = database.hapus_semua_data(user_id)
    await update.message.reply_text(
        f"🗑️ {jumlah_dihapus} transaksi Anda telah dihapus. Saldo kembali ke Rp 0."
    )


async def daftar_transaksi(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    """Command /daftar — tampilkan SEMUA transaksi beserta ID-nya, supaya bisa dipakai di /edit atau /hapus."""
    user_id = update.effective_user.id
    semua = database.get_semua_transaksi(user_id)

    if not semua:
        await update.message.reply_text(
            "📋 Belum ada transaksi tercatat.\n\nCatat dulu dengan `/masuk` atau `/keluar`.",
            parse_mode="Markdown",
        )
        return

    baris = ["📋 *DAFTAR SEMUA TRANSAKSI*\n"]
    for t in semua:
        ikon = "🟢" if t["jenis"] == "masuk" else "🔴"
        tanggal = t["waktu"][:16].replace("T", " ")
        ket = t["keterangan"] or "-"
        baris.append(f"`#{t['id']}` {ikon} {tanggal} — {format_rupiah(t['jumlah'])} ({ket})")

    baris.append(
        "\nGunakan ID (`#angka`) di atas untuk:\n"
        "`/edit <id> <jumlah_baru> <keterangan_baru>`\n"
        "`/hapus <id>`"
    )

    teks_lengkap = "\n".join(baris)

    # Telegram membatasi ~4096 karakter per pesan; potong kalau terlalu panjang.
    if len(teks_lengkap) > 4000:
        teks_lengkap = teks_lengkap[:3900] + "\n\n... (daftar dipotong, terlalu banyak transaksi)"

    await update.message.reply_text(teks_lengkap, parse_mode="Markdown")


async def hapus_transaksi(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    """Format: /hapus <id>"""
    user_id = update.effective_user.id

    if not context.args or not context.args[0].isdigit():
        await update.message.reply_text(
            "⚠️ Format salah.\n\nContoh: `/hapus 3`\n\n"
            "Cek dulu ID transaksinya lewat /daftar.",
            parse_mode="Markdown",
        )
        return

    transaksi_id = int(context.args[0])
    berhasil = database.hapus_transaksi_by_id(transaksi_id, user_id)

    if berhasil:
        saldo_baru = database.get_saldo(user_id)
        await update.message.reply_text(
            f"🗑️ Transaksi `#{transaksi_id}` berhasil dihapus.\n\n"
            f"💰 Saldo sekarang: {format_rupiah(saldo_baru)}",
            parse_mode="Markdown",
        )
    else:
        await update.message.reply_text(
            f"⚠️ Transaksi `#{transaksi_id}` tidak ditemukan, atau bukan milik Anda.\n\n"
            "Cek ID yang benar lewat /daftar.",
            parse_mode="Markdown",
        )


async def edit_transaksi(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    """Format: /edit <id> <jumlah_baru> <keterangan_baru> #KategoriBaru (tag kategori opsional)"""
    user_id = update.effective_user.id

    if len(context.args) < 2 or not context.args[0].isdigit() or not context.args[1].isdigit():
        await update.message.reply_text(
            "⚠️ Format salah.\n\nContoh: `/edit 3 75000 Gaji plus bonus #Gaji`\n\n"
            "Format: `/edit <id> <jumlah_baru> <keterangan_baru> #KategoriBaru`\n"
            "Bagian `#KategoriBaru` opsional. Cek dulu ID transaksinya lewat /daftar.",
            parse_mode="Markdown",
        )
        return

    transaksi_id = int(context.args[0])
    jumlah_baru = int(context.args[1])

    kategori_baru, sisa_args = _ekstrak_kategori(context.args[2:])
    keterangan_baru = " ".join(sisa_args) if sisa_args else None

    berhasil = database.update_transaksi(
        transaksi_id, user_id, jumlah=jumlah_baru, keterangan=keterangan_baru, kategori=kategori_baru
    )

    if berhasil:
        data_terbaru = database.get_transaksi_by_id(transaksi_id, user_id)
        saldo_baru = database.get_saldo(user_id)
        await update.message.reply_text(
            f"✏️ Transaksi `#{transaksi_id}` berhasil diubah.\n\n"
            f"Jumlah baru: {format_rupiah(data_terbaru['jumlah'])}\n"
            f"Kategori: {data_terbaru['kategori']}\n"
            f"Keterangan baru: {data_terbaru['keterangan']}\n\n"
            f"💰 Saldo sekarang: {format_rupiah(saldo_baru)}",
            parse_mode="Markdown",
        )
    else:
        await update.message.reply_text(
            f"⚠️ Transaksi `#{transaksi_id}` tidak ditemukan, atau bukan milik Anda.\n\n"
            "Cek ID yang benar lewat /daftar.",
            parse_mode="Markdown",
        )


async def laporan_kategori(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    """Command /kategori — ringkasan pemasukan & pengeluaran per kategori, 30 hari terakhir."""
    user_id = update.effective_user.id
    data = database.get_laporan_per_kategori(user_id, hari=30)

    if not data:
        await update.message.reply_text(
            "📂 Belum ada transaksi dalam 30 hari terakhir untuk dibuat laporan kategori.\n\n"
            "Catat transaksi dengan tag kategori, contoh:\n`/keluar 20000 Makan siang #Makanan`",
            parse_mode="Markdown",
        )
        return

    baris = ["📂 *LAPORAN PER KATEGORI (30 HARI TERAKHIR)*\n"]
    for row in data:
        net = row["total_masuk"] - row["total_keluar"]
        baris.append(
            f"*{row['kategori']}* ({row['jumlah_transaksi']}x)\n"
            f"  🟢 Masuk: {format_rupiah(row['total_masuk'])}  |  🔴 Keluar: {format_rupiah(row['total_keluar'])}\n"
            f"  Net: {format_rupiah(net)}"
        )

    baris.append(
        "\nTips: tambahkan `#NamaKategori` di akhir command /masuk atau /keluar "
        "untuk mengelompokkan transaksi, contoh:\n`/masuk 50000 Freelance #Sampingan`"
    )

    await update.message.reply_text("\n\n".join(baris), parse_mode="Markdown")


async def export_data(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    """Command /export — kirim semua transaksi sebagai file Excel (.xlsx) yang bisa didownload."""
    user_id = update.effective_user.id
    semua = database.get_semua_transaksi(user_id)

    if not semua:
        await update.message.reply_text(
            "📤 Belum ada transaksi untuk di-export.\n\nCatat dulu dengan /masuk atau /keluar."
        )
        return

    await update.message.reply_text("⏳ Menyiapkan file Excel...")

    path_file = buat_file_export(user_id, semua)
    try:
        with open(path_file, "rb") as f:
            await update.message.reply_document(
                document=f,
                filename=os.path.basename(path_file),
                caption=f"📤 Export {len(semua)} transaksi Anda.",
            )
    finally:
        if os.path.exists(path_file):
            os.remove(path_file)



# HANDLER TOMBOL (CALLBACK QUERY)
# ============================================================

async def button_handler(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    query = update.callback_query
    await query.answer()

    data = query.data

    # --- Navigasi utama ---
    if data == "back_menu":
        await query.edit_message_text(
            MENU_TEXT,
            parse_mode="Markdown",
            reply_markup=main_menu_keyboard(),
        )
        return

    if data == "menu_info":
        await query.edit_message_text(
            INFO_MENU_TEXT,
            parse_mode="Markdown",
            reply_markup=info_menu_keyboard(),
        )
        return

    # --- Detail konten Info & Referensi (isi asli, bukan placeholder) ---
    if data in INFO_CONTENT:
        content = INFO_CONTENT[data]
        await query.edit_message_text(
            content["text"],
            parse_mode="Markdown",
            reply_markup=info_detail_keyboard(),
        )
        return

    # --- Menu fitur lain (sekarang terhubung ke database asli) ---
    user_id = query.from_user.id

    if data == "menu_laporan":
        ringkasan = database.get_ringkasan_laporan(user_id, hari=30)
        saldo = ringkasan["total_masuk"] - ringkasan["total_keluar"]
        text = (
            "📊 *LAPORAN (30 HARI TERAKHIR)*\n\n"
            f"🟢 Total Pemasukan : {format_rupiah(ringkasan['total_masuk'])}\n"
            f"🔴 Total Pengeluaran : {format_rupiah(ringkasan['total_keluar'])}\n"
            f"📈 Selisih (Net) : {format_rupiah(saldo)}\n"
            f"🧾 Jumlah Transaksi : {ringkasan['jumlah_transaksi']} transaksi\n\n"
            "Gunakan /masuk atau /keluar untuk mencatat transaksi baru.\n"
            "Lihat rincian per kategori dengan /kategori, atau download semua data dengan /export."
        )
    elif data == "menu_saldo":
        saldo = database.get_saldo(user_id)
        text = (
            "💰 *SALDO ANDA SAAT INI*\n\n"
            f"{format_rupiah(saldo)}\n\n"
            "Saldo dihitung dari seluruh riwayat pemasukan dikurangi pengeluaran yang tercatat."
        )
    elif data == "menu_transaksi":
        daftar = database.get_transaksi_terakhir(user_id, limit=5)
        if not daftar:
            text = (
                "🧾 *TRANSAKSI TERAKHIR*\n\n"
                "Belum ada transaksi tercatat.\n\n"
                "Catat transaksi pertama Anda dengan:\n"
                "`/masuk 50000 Gaji #Gaji`\n"
                "`/keluar 20000 Makan siang #Makanan`"
            )
        else:
            baris = []
            for t in daftar:
                ikon = "🟢" if t["jenis"] == "masuk" else "🔴"
                tanggal = t["waktu"][:16].replace("T", " ")
                ket = t["keterangan"] or "-"
                baris.append(f"{ikon} {tanggal} — {format_rupiah(t['jumlah'])} [{t['kategori']}] ({ket})")
            text = (
                "🧾 *5 TRANSAKSI TERAKHIR*\n\n" + "\n".join(baris) +
                "\n\nCatat transaksi baru dengan /masuk atau /keluar. Lihat semua dengan /daftar."
            )
    elif data == "menu_data":
        total = database.get_jumlah_data(user_id)
        text = (
            "📁 *DATA*\n\n"
            f"Total transaksi tersimpan: *{total}* baris\n\n"
            "Command terkait data:\n"
            "`/masuk <jumlah> <ket> #Kategori` — catat pemasukan\n"
            "`/keluar <jumlah> <ket> #Kategori` — catat pengeluaran\n"
            "`/daftar` — lihat semua transaksi beserta ID-nya\n"
            "`/edit <id> <jumlah> <ket> #Kategori` — ubah satu transaksi\n"
            "`/hapus <id>` — hapus satu transaksi\n"
            "`/kategori` — laporan ringkasan per kategori\n"
            "`/export` — download semua data sebagai file Excel\n"
            "`/reset_data` — hapus SEMUA data (tidak bisa dibatalkan)"
        )
    elif data == "menu_help":
        await query.edit_message_text(
            "❓ *BANTUAN*\n\nPilih bantuan yang Anda perlukan.",
            parse_mode="Markdown",
            reply_markup=help_keyboard(),
        )
        return
    elif data == "help_cara":
        text = (
            "📖 *CARA MENGGUNAKAN BOT*\n\n"
            "1. Ketik /start untuk memulai.\n"
            "2. Ketik /menu untuk membuka menu.\n"
            "3. Ketik /info untuk membuka referensi langsung.\n"
            "4. Tekan tombol sesuai kebutuhan.\n"
            "5. Gunakan tombol 🏠 Menu Utama untuk kembali.\n\n"
            "Pesan teks biasa tidak diproses oleh bot ini."
        )
        await query.edit_message_text(
            text,
            parse_mode="Markdown",
            reply_markup=back_menu_keyboard(),
        )
        return
    elif data == "help_command":
        text = (
            "📋 *DAFTAR COMMAND*\n\n"
            "/start — Memulai bot\n"
            "/menu — Menu utama\n"
            "/info — Menu info & referensi\n"
            "/help — Bantuan penggunaan"
        )
        await query.edit_message_text(
            text,
            parse_mode="Markdown",
            reply_markup=back_menu_keyboard(),
        )
        return
    else:
        text = "Perintah tombol tidak dikenal. Silakan kembali ke menu utama."

    await query.edit_message_text(
        text,
        parse_mode="Markdown",
        reply_markup=back_menu_keyboard(),
    )


# ============================================================
# MAIN
# ============================================================

def main() -> None:
    if not TOKEN:
        raise RuntimeError(
            "BOT_TOKEN belum ditemukan. Atur environment variable BOT_TOKEN terlebih dahulu."
        )

    database.init_db()

    application = ApplicationBuilder().token(TOKEN).build()

    application.add_handler(CommandHandler("start", start))
    application.add_handler(CommandHandler("help", help_command))
    application.add_handler(CommandHandler("menu", menu))
    application.add_handler(CommandHandler("info", info_command))
    application.add_handler(CommandHandler("masuk", catat_masuk))
    application.add_handler(CommandHandler("keluar", catat_keluar))
    application.add_handler(CommandHandler("reset_data", reset_data))
    application.add_handler(CommandHandler("daftar", daftar_transaksi))
    application.add_handler(CommandHandler("hapus", hapus_transaksi))
    application.add_handler(CommandHandler("edit", edit_transaksi))
    application.add_handler(CommandHandler("kategori", laporan_kategori))
    application.add_handler(CommandHandler("export", export_data))
    application.add_handler(CallbackQueryHandler(button_handler))

    application.run_polling()


if __name__ == "__main__":
    main()
